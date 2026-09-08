import io
import os
import json
import time
import base64
import mammoth
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from django.shortcuts import render, redirect
from django.urls import reverse
from django.http import JsonResponse, HttpResponse, StreamingHttpResponse, Http404
from django.contrib.auth import authenticate, login as auth_login, logout as auth_logout, get_user_model
from django.views.decorators.csrf import csrf_exempt
from . import mongo
from . import otp_service

# ================= AUTHENTICATION & 2-STEP VERIFICATION VIEWS =================

def login_view(request):
    if request.user.is_authenticated and request.user.is_staff:
        return redirect("/")
        
    error = None
    if request.method == "POST":
        username = request.POST.get("username", "").strip()
        password = request.POST.get("password", "").strip()
        
        user = authenticate(request, username=username, password=password)
        if user is not None:
            if user.is_staff:
                # Initiate 2-Step Verification
                phone_number = os.getenv("ADMIN_PHONE_NUMBER", "+919876543210").strip()
                otp_code = otp_service.generate_otp()
                
                request.session['pre_2fa_user_id'] = user.id
                request.session['otp_code'] = otp_code
                request.session['otp_expires_at'] = time.time() + 300  # 5 minutes validity
                request.session['otp_attempts'] = 0
                request.session['otp_phone'] = phone_number
                request.session['otp_sent_at'] = time.time()
                
                # Send SMS via gateway / console
                otp_service.send_otp_sms(phone_number, otp_code)
                
                return redirect(reverse("documents:verify_otp"))
            else:
                error = "Access denied. Only administrators can log into the management console."
        else:
            error = "Invalid username or password. Please try again."
            
    return render(request, "documents/login.html", {"error": error})

def verify_otp_view(request):
    if request.user.is_authenticated and request.user.is_staff:
        return redirect("/")
        
    user_id = request.session.get('pre_2fa_user_id')
    if not user_id:
        return redirect(reverse("documents:login"))
        
    phone_number = request.session.get('otp_phone', '+91 ••••• ••123')
    masked_phone = otp_service.mask_phone_number(phone_number)
    error = None
    
    if request.method == "POST":
        submitted_otp = request.POST.get("otp", "").strip()
        actual_otp = request.session.get('otp_code')
        expires_at = request.session.get('otp_expires_at', 0)
        attempts = request.session.get('otp_attempts', 0)
        
        # Check attempts limit
        if attempts >= 3:
            # Clear 2FA session on security lockout
            for key in ['pre_2fa_user_id', 'otp_code', 'otp_expires_at', 'otp_attempts', 'otp_phone', 'otp_sent_at']:
                request.session.pop(key, None)
            return render(request, "documents/login.html", {
                "error": "Account locked after 3 failed OTP attempts. Please log in again."
            })
            
        # Check expiration
        if time.time() > expires_at:
            error = "Verification code has expired. Please click 'Resend OTP' below."
        elif submitted_otp == actual_otp:
            # Successful 2FA verification!
            User = get_user_model()
            try:
                user = User.objects.get(id=user_id)
                auth_login(request, user)
                
                # Clean up session
                for key in ['pre_2fa_user_id', 'otp_code', 'otp_expires_at', 'otp_attempts', 'otp_phone', 'otp_sent_at']:
                    request.session.pop(key, None)
                    
                return redirect("/")
            except User.DoesNotExist:
                return redirect(reverse("documents:login"))
        else:
            attempts += 1
            request.session['otp_attempts'] = attempts
            remaining = 3 - attempts
            if remaining <= 0:
                for key in ['pre_2fa_user_id', 'otp_code', 'otp_expires_at', 'otp_attempts', 'otp_phone', 'otp_sent_at']:
                    request.session.pop(key, None)
                return render(request, "documents/login.html", {
                    "error": "Maximum OTP attempts exceeded. Please log in again."
                })
            else:
                error = f"Invalid verification code. {remaining} attempt(s) remaining."
                
    return render(request, "documents/verify_otp.html", {
        "masked_phone": masked_phone,
        "error": error
    })

def resend_otp_view(request):
    user_id = request.session.get('pre_2fa_user_id')
    if not user_id:
        return redirect(reverse("documents:login"))
        
    last_sent = request.session.get('otp_sent_at', 0)
    # 60s cooldown
    if time.time() - last_sent >= 60:
        phone_number = request.session.get('otp_phone', os.getenv("ADMIN_PHONE_NUMBER", "+919876543210").strip())
        new_otp = otp_service.generate_otp()
        request.session['otp_code'] = new_otp
        request.session['otp_expires_at'] = time.time() + 300
        request.session['otp_attempts'] = 0
        request.session['otp_sent_at'] = time.time()
        
        otp_service.send_otp_sms(phone_number, new_otp)
        
    return redirect(reverse("documents:verify_otp"))

def logout_view(request):
    auth_logout(request)
    return redirect("/")

@csrf_exempt
def api_auth_login_view(request):
    if request.method != "POST":
        return JsonResponse({"success": False, "error": "POST method required"}, status=405)
    
    username = ""
    password = ""
    
    # Support both JSON body and standard Form POST
    if request.content_type == "application/json" and request.body:
        try:
            body_data = json.loads(request.body)
            username = body_data.get("username", "").strip()
            password = body_data.get("password", "").strip()
        except Exception:
            pass
    
    if not username:
        username = request.POST.get("username", "").strip()
        password = request.POST.get("password", "").strip()
        
    if not username or not password:
        return JsonResponse({"success": False, "error": "Username and password are required"}, status=400)
        
    user = authenticate(request, username=username, password=password)
    if user is not None:
        if user.is_staff:
            # 2FA initialization for mobile/REST
            phone_number = os.getenv("ADMIN_PHONE_NUMBER", "+919876543210").strip()
            otp_code = otp_service.generate_otp()
            
            request.session['pre_2fa_user_id'] = user.id
            request.session['otp_code'] = otp_code
            request.session['otp_expires_at'] = time.time() + 300
            request.session['otp_attempts'] = 0
            request.session['otp_phone'] = phone_number
            request.session['otp_sent_at'] = time.time()
            
            otp_service.send_otp_sms(phone_number, otp_code)
            
            return JsonResponse({
                "success": True,
                "step": "otp_required",
                "phone_masked": otp_service.mask_phone_number(phone_number),
                "message": f"Verification code sent to {otp_service.mask_phone_number(phone_number)}"
            })
        else:
            return JsonResponse({
                "success": False, 
                "error": "Access denied. Only administrators have management privileges."
            }, status=403)
    else:
        return JsonResponse({"success": False, "error": "Invalid username or password"}, status=401)

@csrf_exempt
def api_auth_verify_otp_view(request):
    if request.method != "POST":
        return JsonResponse({"success": False, "error": "POST method required"}, status=405)
        
    user_id = request.session.get('pre_2fa_user_id')
    if not user_id:
        return JsonResponse({"success": False, "error": "No pending 2FA login session found. Please log in again."}, status=400)
        
    otp = ""
    if request.content_type == "application/json" and request.body:
        try:
            body_data = json.loads(request.body)
            otp = str(body_data.get("otp", "")).strip()
        except Exception:
            pass
    if not otp:
        otp = request.POST.get("otp", "").strip()
        
    if not otp:
        return JsonResponse({"success": False, "error": "OTP is required"}, status=400)
        
    actual_otp = request.session.get('otp_code')
    expires_at = request.session.get('otp_expires_at', 0)
    attempts = request.session.get('otp_attempts', 0)
    
    if attempts >= 3:
        for key in ['pre_2fa_user_id', 'otp_code', 'otp_expires_at', 'otp_attempts', 'otp_phone', 'otp_sent_at']:
            request.session.pop(key, None)
        return JsonResponse({"success": False, "error": "Maximum OTP attempts exceeded. Please log in again."}, status=403)
        
    if time.time() > expires_at:
        return JsonResponse({"success": False, "error": "OTP has expired. Please request a new one."}, status=400)
        
    if otp == actual_otp:
        User = get_user_model()
        try:
            user = User.objects.get(id=user_id)
            auth_login(request, user)
            for key in ['pre_2fa_user_id', 'otp_code', 'otp_expires_at', 'otp_attempts', 'otp_phone', 'otp_sent_at']:
                request.session.pop(key, None)
            return JsonResponse({
                "success": True,
                "is_authenticated": True,
                "is_admin": True,
                "username": user.username,
                "message": f"2-Step verification successful! Welcome, {user.username}."
            })
        except User.DoesNotExist:
            return JsonResponse({"success": False, "error": "User not found"}, status=404)
    else:
        attempts += 1
        request.session['otp_attempts'] = attempts
        remaining = 3 - attempts
        if remaining <= 0:
            for key in ['pre_2fa_user_id', 'otp_code', 'otp_expires_at', 'otp_attempts', 'otp_phone', 'otp_sent_at']:
                request.session.pop(key, None)
            return JsonResponse({"success": False, "error": "Maximum OTP attempts exceeded. Please log in again."}, status=403)
        return JsonResponse({"success": False, "error": f"Invalid verification code. {remaining} attempt(s) remaining."}, status=400)

@csrf_exempt
def api_auth_resend_otp_view(request):
    user_id = request.session.get('pre_2fa_user_id')
    if not user_id:
        return JsonResponse({"success": False, "error": "No pending 2FA login session found."}, status=400)
        
    last_sent = request.session.get('otp_sent_at', 0)
    if time.time() - last_sent < 60:
        wait = int(60 - (time.time() - last_sent))
        return JsonResponse({"success": False, "error": f"Please wait {wait} seconds before requesting a new OTP."}, status=429)
        
    phone_number = request.session.get('otp_phone', os.getenv("ADMIN_PHONE_NUMBER", "+919876543210").strip())
    new_otp = otp_service.generate_otp()
    request.session['otp_code'] = new_otp
    request.session['otp_expires_at'] = time.time() + 300
    request.session['otp_attempts'] = 0
    request.session['otp_sent_at'] = time.time()
    
    otp_service.send_otp_sms(phone_number, new_otp)
    return JsonResponse({"success": True, "message": f"New OTP sent to {otp_service.mask_phone_number(phone_number)}"})

def api_auth_status_view(request):
    is_auth = request.user.is_authenticated
    is_staff = is_auth and request.user.is_staff
    return JsonResponse({
        "success": True,
        "is_authenticated": is_auth,
        "is_admin": is_staff,
        "username": request.user.username if is_auth else ""
    })

@csrf_exempt
def api_auth_logout_view(request):
    auth_logout(request)
    return JsonResponse({"success": True, "message": "Successfully logged out"})

def pwa_manifest_view(request):
    manifest = {
        "name": "Subject Bank - Academic Vault",
        "short_name": "SubjectBank",
        "description": "Store and preview PDF, PPT, Word, and Image documents stored in MongoDB.",
        "start_url": "/",
        "display": "standalone",
        "background_color": "#ffffff",
        "theme_color": "#4f46e5",
        "icons": [
            {
                "src": "/static/documents/icon-192.png",
                "sizes": "192x192",
                "type": "image/png"
            },
            {
                "src": "/static/documents/icon-512.png",
                "sizes": "512x512",
                "type": "image/png"
            }
        ]
    }
    return JsonResponse(manifest)

def health_check_view(request):
    """Ultra-fast keep-alive endpoint for cron-job.org / UptimeRobot to keep Render awake 24/7."""
    return HttpResponse("OK", content_type="text/plain", status=200)


# ================= PAGE VIEWS =================

def index_view(request):
    stats = mongo.get_stats()
    subjects = mongo.get_all_subjects_with_stats()
    return render(request, "documents/index.html", {
        "stats": stats,
        "subjects": subjects,
        "is_admin": request.user.is_authenticated and request.user.is_staff
    })

# ================= SUBJECT & FOLDER API =================

def api_subjects_view(request):
    subjects = mongo.get_all_subjects_with_stats()
    stats = mongo.get_stats()
    return JsonResponse({
        "success": True,
        "subjects": subjects,
        "stats": stats,
        "is_admin": request.user.is_authenticated and request.user.is_staff
    })

@csrf_exempt
def api_create_subject_view(request):
    if not (request.user.is_authenticated and request.user.is_staff):
        return JsonResponse({"success": False, "error": "Unauthorized. Only administrators can create subjects."}, status=403)
        
    if request.method != "POST":
        return JsonResponse({"success": False, "error": "POST method required"}, status=405)
        
    name = request.POST.get("name", "").strip()
    description = request.POST.get("description", "").strip()
    
    try:
        mongo.create_subject(name=name, description=description)
        return JsonResponse({"success": True, "message": f"Subject '{name}' created successfully!"})
    except Exception as e:
        return JsonResponse({"success": False, "error": str(e)}, status=400)

@csrf_exempt
def api_delete_subject_view(request, subject_name):
    if not (request.user.is_authenticated and request.user.is_staff):
        return JsonResponse({"success": False, "error": "Unauthorized. Only administrators can delete subjects."}, status=403)
        
    if request.method not in ["POST", "DELETE"]:
        return JsonResponse({"success": False, "error": "Invalid method"}, status=405)
        
    try:
        mongo.delete_subject(subject_name)
        return JsonResponse({"success": True, "message": f"Subject '{subject_name}' deleted."})
    except Exception as e:
        return JsonResponse({"success": False, "error": str(e)}, status=400)

@csrf_exempt
def api_create_folder_view(request):
    if not (request.user.is_authenticated and request.user.is_staff):
        return JsonResponse({"success": False, "error": "Unauthorized. Only administrators can create folders."}, status=403)
        
    if request.method != "POST":
        return JsonResponse({"success": False, "error": "POST method required"}, status=405)
        
    subject = request.POST.get("subject", "").strip()
    folder_name = request.POST.get("folder_name", "").strip()
    
    try:
        mongo.create_folder(subject_name=subject, folder_name=folder_name)
        return JsonResponse({"success": True, "message": f"Folder '{folder_name}' created in '{subject}'!"})
    except Exception as e:
        return JsonResponse({"success": False, "error": str(e)}, status=400)

@csrf_exempt
def api_delete_folder_view(request):
    if not (request.user.is_authenticated and request.user.is_staff):
        return JsonResponse({"success": False, "error": "Unauthorized. Only administrators can delete folders."}, status=403)
        
    if request.method not in ["POST", "DELETE"]:
        return JsonResponse({"success": False, "error": "Invalid method"}, status=405)
        
    subject = request.POST.get("subject", "").strip()
    folder_name = request.POST.get("folder_name", "").strip()
    
    try:
        mongo.delete_folder(subject_name=subject, folder_name=folder_name)
        return JsonResponse({"success": True, "message": f"Folder '{folder_name}' deleted."})
    except Exception as e:
        return JsonResponse({"success": False, "error": str(e)}, status=400)

# ================= DOCUMENT API =================

def api_documents_view(request):
    query = request.GET.get("q", "").strip()
    subject = request.GET.get("subject", "").strip()
    folder = request.GET.get("folder", "").strip()
    category = request.GET.get("category", "").strip()
    
    docs = mongo.list_documents(query=query, subject=subject, folder=folder, category=category)
    stats = mongo.get_stats()
    
    return JsonResponse({
        "success": True,
        "count": len(docs),
        "documents": docs,
        "stats": stats,
        "is_admin": request.user.is_authenticated and request.user.is_staff
    })

@csrf_exempt
def api_upload_view(request):
    # Enforce Admin Check
    if not (request.user.is_authenticated and request.user.is_staff):
        return JsonResponse({
            "success": False, 
            "error": "Access Denied: Only authenticated administrators can upload documents."
        }, status=403)
    
    if request.method != "POST":
        return JsonResponse({"success": False, "error": "POST method required"}, status=405)
    
    # Support both single file ('file') and multiple files ('files' or 'file' multiple times)
    files = request.FILES.getlist("files")
    if not files:
        files = request.FILES.getlist("file")
        
    if not files:
        return JsonResponse({"success": False, "error": "No file(s) selected for upload"}, status=400)
    
    subject = request.POST.get("subject", "").strip() or "General"
    folder = request.POST.get("folder", "").strip() or "General"
    tags = request.POST.get("tags", "")
    batch_title = request.POST.get("title", "").strip()
    
    saved_ids = []
    saved_names = []
    
    try:
        for idx, file_obj in enumerate(files):
            # If a single file and a title was given, use it. Otherwise use the file's original name.
            if len(files) == 1 and batch_title:
                doc_title = batch_title
            elif batch_title and len(files) > 1:
                # If a batch prefix was provided, format as "Prefix - Filename"
                doc_title = f"{batch_title} - {file_obj.name}"
            else:
                doc_title = file_obj.name
                
            file_id = mongo.save_document(
                file_obj=file_obj,
                title=doc_title,
                subject=subject,
                folder=folder,
                tags=tags
            )
            saved_ids.append(file_id)
            saved_names.append(file_obj.name)
            
        mongo.invalidate_cache()
        
        msg = f"Successfully stored {len(saved_ids)} document(s) in {subject} / {folder}!"
        return JsonResponse({
            "success": True,
            "count": len(saved_ids),
            "id": saved_ids[0] if saved_ids else None,
            "ids": saved_ids,
            "filenames": saved_names,
            "message": msg,
        })
    except Exception as e:
        return JsonResponse({"success": False, "error": str(e)}, status=500)

@csrf_exempt
def api_delete_view(request, file_id):
    # Enforce Admin Check
    if not (request.user.is_authenticated and request.user.is_staff):
        return JsonResponse({
            "success": False, 
            "error": "Access Denied: Only authenticated administrators can delete documents."
        }, status=403)

    if request.method not in ["POST", "DELETE"]:
        return JsonResponse({"success": False, "error": "Invalid request method"}, status=405)
        
    deleted = mongo.delete_document(file_id)
    if deleted:
        return JsonResponse({"success": True, "message": "Document deleted from MongoDB"})
    else:
        return JsonResponse({"success": False, "error": "Document not found or could not be deleted"}, status=404)

# ================= STREAMING & PREVIEWS =================

def _stream_gridfs_file(grid_out, chunk_size=256 * 1024):
    while True:
        chunk = grid_out.read(chunk_size)
        if not chunk:
            break
        yield chunk

def view_file_inline_view(request, file_id):
    grid_out = mongo.get_document(file_id)
    if not grid_out:
        raise Http404("Document not found")
        
    content_type = getattr(grid_out, 'contentType', None) or getattr(grid_out, 'content_type', 'application/octet-stream')
    filename = grid_out.filename or "document"
    
    response = StreamingHttpResponse(
        _stream_gridfs_file(grid_out),
        content_type=content_type
    )
    response["Content-Disposition"] = f'inline; filename="{filename}"'
    response["Content-Length"] = grid_out.length
    return response

def download_file_view(request, file_id):
    grid_out = mongo.get_document(file_id)
    if not grid_out:
        raise Http404("Document not found")
        
    content_type = getattr(grid_out, 'contentType', None) or getattr(grid_out, 'content_type', 'application/octet-stream')
    filename = grid_out.filename or "document"
    
    response = StreamingHttpResponse(
        _stream_gridfs_file(grid_out),
        content_type=content_type
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    response["Content-Length"] = grid_out.length
    return response

def _extract_pptx_slides(file_bytes):
    prs = Presentation(io.BytesIO(file_bytes))
    slides_data = []
    
    for idx, slide in enumerate(prs.slides, 1):
        slide_info = {
            "slide_number": idx,
            "title": "",
            "content": [],
            "images": [],
            "tables": []
        }
        
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text.strip()
                if not text:
                    continue
                if shape == slide.shapes.title:
                    slide_info["title"] = text
                else:
                    paragraphs = []
                    for p in shape.text_frame.paragraphs:
                        p_text = p.text.strip()
                        if p_text:
                            paragraphs.append({
                                "text": p_text,
                                "level": p.level
                            })
                    if paragraphs:
                        slide_info["content"].append(paragraphs)
                        
            if shape.has_table:
                table_data = []
                for row in shape.table.rows:
                    row_data = [cell.text.strip() for cell in row.cells]
                    table_data.append(row_data)
                slide_info["tables"].append(table_data)
                
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                img_bytes = shape.image.blob
                ct = shape.image.content_type
                b64 = base64.b64encode(img_bytes).decode('utf-8')
                slide_info["images"].append(f"data:{ct};base64,{b64}")
                
        if not slide_info["title"]:
            slide_info["title"] = f"Slide {idx}"
            
        slides_data.append(slide_info)
    return slides_data

def api_preview_content_view(request, file_id):
    grid_out = mongo.get_document(file_id)
    if not grid_out:
        return JsonResponse({"success": False, "error": "Document not found"}, status=404)
        
    filename = (grid_out.filename or "").lower()
    title = getattr(grid_out, "title", grid_out.filename)
    category = getattr(grid_out, "category", "other")
    
    view_url = f"/documents/{file_id}/view/"
    download_url = f"/documents/{file_id}/download/"

    if category == "pdf" or filename.endswith(".pdf"):
        return JsonResponse({
            "success": True,
            "type": "pdf",
            "url": view_url,
            "title": title,
            "filename": grid_out.filename
        })
        
    if category == "image" or filename.endswith((".png", ".jpg", ".jpeg", ".gif", ".webp", ".svg")):
        return JsonResponse({
            "success": True,
            "type": "image",
            "url": view_url,
            "title": title,
            "filename": grid_out.filename
        })

    # Read binary bytes from GridFS
    file_bytes = grid_out.read()

    # Word Documents
    if filename.endswith(".docx"):
        try:
            result = mammoth.convert_to_html(io.BytesIO(file_bytes))
            return JsonResponse({
                "success": True,
                "type": "docx",
                "html": result.value,
                "title": title,
                "filename": grid_out.filename
            })
        except Exception as e:
            return JsonResponse({
                "success": False,
                "type": "unsupported",
                "error": f"Error parsing Word document: {str(e)}",
                "download_url": download_url
            })

    # PowerPoint Presentations
    if filename.endswith(".pptx"):
        try:
            slides = _extract_pptx_slides(file_bytes)
            return JsonResponse({
                "success": True,
                "type": "pptx",
                "slides": slides,
                "title": title,
                "filename": grid_out.filename
            })
        except Exception as e:
            return JsonResponse({
                "success": False,
                "type": "unsupported",
                "error": f"Error parsing presentation: {str(e)}",
                "download_url": download_url
            })

    # Plain text / code files
    if filename.endswith((".txt", ".md", ".csv", ".json", ".log", ".py", ".js", ".html")):
        try:
            text_content = file_bytes.decode('utf-8', errors='replace')
            return JsonResponse({
                "success": True,
                "type": "text",
                "text": text_content,
                "title": title,
                "filename": grid_out.filename
            })
        except Exception:
            pass

    return JsonResponse({
        "success": False,
        "type": "unsupported",
        "error": "This file format cannot be parsed directly in the browser.",
        "download_url": download_url,
        "title": title,
        "filename": grid_out.filename
    })
