import os
import io
import django
from django.core.files.uploadedfile import SimpleUploadedFile

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'subject_bank.settings')
django.setup()

from django.test import Client
from django.contrib.auth.models import User
from documents import mongo
from docx import Document
from pptx import Presentation

def run_tests():
    print("=== STARTING FULL ROLE-BASED ACCESS & SUBJECT HIERARCHY TESTS ===")
    
    viewer_client = Client()
    admin_client = Client()
    
    # 1. Verify Unauthenticated Viewer Restrictions
    print("\n--- 1. Testing Viewer (Public / Non-Admin) Restrictions ---")
    r_home = viewer_client.get('/')
    assert r_home.status_code == 200
    assert "Admin Login" in r_home.content.decode('utf-8')
    print("  [OK] Viewer accesses homepage without admin privileges")
    
    r_upload_fail = viewer_client.post('/api/documents/upload/', {'title': 'Hacker Note'})
    assert r_upload_fail.status_code == 403
    print("  [OK] Viewer blocked from uploading documents (403 Forbidden)")
    
    r_create_subj_fail = viewer_client.post('/api/subjects/create/', {'name': 'Hacker Subject'})
    assert r_create_subj_fail.status_code == 403
    print("  [OK] Viewer blocked from creating subjects (403 Forbidden)")
    
    r_create_folder_fail = viewer_client.post('/api/folders/create/', {'subject': 'Cloud Computing', 'folder_name': 'Hacker Folder'})
    assert r_create_folder_fail.status_code == 403
    print("  [OK] Viewer blocked from creating folders (403 Forbidden)")

    # 2. Admin Authentication
    print("\n--- 2. Testing Admin Login ---")
    from django.contrib.auth import get_user_model
    User = get_user_model()
    test_admin, _ = User.objects.get_or_create(username='testadmin', defaults={'email': 'test@example.com', 'is_staff': True, 'is_superuser': True})
    test_admin.set_password('testpass123')
    test_admin.is_staff = True
    test_admin.save()

    login_res = admin_client.post('/login/', {'username': 'testadmin', 'password': 'testpass123'})
    assert login_res.status_code == 302 # redirect to /verify-otp/
    assert '/verify-otp/' in login_res.url
    
    # Extract OTP from session and verify
    otp_code = admin_client.session.get('otp_code')
    assert otp_code is not None and len(otp_code) == 6
    r_otp = admin_client.post('/verify-otp/', {'otp': otp_code})
    assert r_otp.status_code == 302
    assert r_otp.url == '/'

    r_admin_home = admin_client.get('/')
    assert "Logout" in r_admin_home.content.decode('utf-8')
    print("  [OK] Admin successfully authenticated with 2-Step Verification (Password + OTP)")


    # 3. Admin creates Subject & Folder
    print("\n--- 3. Testing Subject & Folder Creation by Admin ---")
    subj_name = "Cloud Computing Core"
    # Ensure clean state by deleting if already exists
    admin_client.post(f'/api/subjects/{subj_name}/delete/')
    
    r_subj = admin_client.post('/api/subjects/create/', {
        'name': subj_name,
        'description': 'Virtualization, Cloud Storage, and Distributed Systems'
    })
    assert r_subj.status_code == 200
    print(f"  [OK] Admin created subject '{subj_name}'")
    
    folder_name = "Unit 1 - Cloud Architecture"
    r_folder = admin_client.post('/api/folders/create/', {
        'subject': subj_name,
        'folder_name': folder_name
    })
    assert r_folder.status_code == 200
    print(f"  [OK] Admin created folder '{folder_name}' in '{subj_name}'")

    # 4. Admin Uploads Material to Subject & Folder
    print("\n--- 4. Testing Admin File Upload into Specific Subject & Folder ---")
    doc = Document()
    doc.add_heading('Cloud Architecture Overview', 0)
    doc.add_paragraph('Infrastructure as a Service (IaaS), Platform as a Service (PaaS), and Software as a Service (SaaS).')
    buf = io.BytesIO()
    doc.save(buf)
    
    file_upload = SimpleUploadedFile("Cloud_Lecture1.docx", buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    
    r_upload = admin_client.post('/api/documents/upload/', {
        'file': file_upload,
        'title': 'Lecture 1: Cloud Service Models',
        'subject': subj_name,
        'folder': folder_name,
        'tags': 'cloud,iaas,paas,saas'
    })
    assert r_upload.status_code == 200
    doc_id = r_upload.json()['id']
    print(f"  [OK] Admin uploaded Word doc into '{subj_name} / {folder_name}' -> ID: {doc_id}")

    # 5. Viewer Browses Subject, Folder & Previews Document Live
    print("\n--- 5. Testing Viewer Accessing Subject & Viewing Document ---")
    r_docs = viewer_client.get(f'/api/documents/?subject={subj_name}&folder={folder_name}')
    assert r_docs.status_code == 200
    docs_data = r_docs.json()
    assert docs_data['count'] >= 1
    print(f"  [OK] Viewer found {docs_data['count']} document(s) in '{subj_name} / {folder_name}'")
    
    # Viewer opens in-browser preview
    r_prev = viewer_client.get(f'/api/documents/{doc_id}/preview-content/')
    assert r_prev.status_code == 200
    prev_data = r_prev.json()
    assert prev_data['type'] == 'docx'
    assert 'Cloud Architecture' in prev_data['html']
    print("  [OK] Viewer previewed Word document live in-browser without downloading!")

    # 6. Admin Deletion
    print("\n--- 6. Testing Admin Deletion Privileges ---")
    r_del = admin_client.post(f'/api/documents/{doc_id}/delete/')
    assert r_del.status_code == 200
    print("  [OK] Admin successfully deleted document from MongoDB GridFS")

    # Clean up subject
    admin_client.post(f'/api/subjects/{subj_name}/delete/')
    print(f"  [OK] Admin cleaned up test subject '{subj_name}'")

    # 7. Mobile REST API Authentication
    print("\n--- 7. Testing Mobile JSON REST Auth APIs ---")
    mobile_client = Client()
    # Anonymous status check
    r_status = mobile_client.get('/api/auth/status/')
    assert r_status.json()['is_authenticated'] == False
    # Mobile login
    r_mob_login = mobile_client.post('/api/auth/login/', {'username': 'testadmin', 'password': 'testpass123'})
    assert r_mob_login.status_code == 200
    assert r_mob_login.json().get('step') == 'otp_required'
    mob_otp = mobile_client.session.get('otp_code')
    assert mob_otp is not None and len(mob_otp) == 6

    # Mobile OTP verify
    r_mob_verify = mobile_client.post('/api/auth/verify-otp/', {'otp': mob_otp}, content_type='application/json')
    assert r_mob_verify.status_code == 200
    assert r_mob_verify.json().get('is_admin') == True

    # Mobile status after login
    r_status2 = mobile_client.get('/api/auth/status/')
    assert r_status2.json()['is_admin'] == True
    # Mobile logout
    r_mob_logout = mobile_client.post('/api/auth/logout/')
    assert r_mob_logout.status_code == 200
    print("  [OK] Mobile JSON REST 2-Step Auth (login, OTP verify, status, logout) verified successfully!")


    # Clean up test admin
    test_admin.delete()
    print("  [OK] Cleaned up temporary test admin user")

    print("\n=== ALL ROLE-BASED ACCESS, SUBJECT HIERARCHY & MOBILE API TESTS PASSED! ===")


if __name__ == '__main__':
    run_tests()

